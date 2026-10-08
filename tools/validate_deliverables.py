from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "智能驾驶汽车用户使用体验调研问卷_1008.docx"
FINAL = ROOT / "智能驾驶汽车用户使用体验调研问卷_规整及系统设计_1008.docx"
SQL = ROOT / "问卷调查系统_数据库管理_MySQL8.sql"
EXPECTED_SOURCE_SHA = "0B9301A49A1BE69AD891706A4D4BEA1764D8431663C72FCA7F8BA85B396EA695"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def part_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {name: hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist()}


assert sha(SOURCE) == EXPECTED_SOURCE_SHA, "Source DOCX changed"
assert FINAL.exists() and FINAL.stat().st_size > 100_000, "Final DOCX missing or unexpectedly small"
assert SQL.exists() and SQL.stat().st_size > 10_000, "SQL deliverable missing or unexpectedly small"

doc = Document(FINAL)
questions = [p.text for p in doc.paragraphs if p.style.name == "SurveyQuestion"]
numbers = [re.match(r"^(Q\d+)", text).group(1) for text in questions]
expected = [f"Q{i:02d}" for i in range(1, 62)]
assert numbers == expected, f"Question sequence mismatch: {numbers}"

all_text = "\n".join(p.text for p in doc.paragraphs)
assert "【留一幅中国地图" not in all_text
assert "【画一幅统计图" not in all_text
assert "【统计25个主要品牌" not in all_text
assert "Q74  您对以下政策" not in all_text
assert "Q01–Q06" in all_text and "Q60–Q61" in all_text
assert "https://ads-experience-survey-ui.shchen-lmars.chatgpt.site/" in all_text
assert "https://ads-experience-survey-ui.shchen-lmars.chatgpt.site/admin.html" in all_text
assert len(doc.tables) == 54, f"Expected 54 tables, found {len(doc.tables)}"
assert len(doc.inline_shapes) == 4, f"Expected 4 figures, found {len(doc.inline_shapes)}"

with zipfile.ZipFile(FINAL) as archive:
    names = archive.namelist()
    assert len([n for n in names if n.startswith("word/media/")]) == 4
    footer = archive.read("word/footer1.xml").decode("utf-8")
    assert "PAGE" in footer
    settings = archive.read("word/settings.xml").decode("utf-8")
    assert "updateFields" in settings

source_parts = part_hashes(SOURCE)
final_parts = part_hashes(FINAL)
preserve = [
    "word/header1.xml",
    "word/footer1.xml",
    "word/theme/theme1.xml",
    "word/numbering.xml",
    "word/styles.xml",
    "word/webSettings.xml",
    "word/fontTable.xml",
    "word/footnotes.xml",
    "word/endnotes.xml",
    "customXml/item1.xml",
    "customXml/itemProps1.xml",
    "customXml/_rels/item1.xml.rels",
]
changed_preserve = [name for name in preserve if source_parts.get(name) != final_parts.get(name)]
assert not changed_preserve, f"Preserve-only package parts changed: {changed_preserve}"

sql = SQL.read_text(encoding="utf-8")
for idx in range(1, 6):
    assert f"reserved_char_{idx:02d} VARCHAR(255)" in sql
for required in ["CREATE TABLE IF NOT EXISTS survey_response", "followup_contact", "admin_audit_log", "CREATE OR REPLACE VIEW v_response_overview"]:
    assert required in sql

print("VALIDATION_OK")
print(f"docx_sha256={sha(FINAL)}")
print(f"docx_paragraphs={len(doc.paragraphs)}")
print(f"docx_tables={len(doc.tables)}")
print(f"docx_figures={len(doc.inline_shapes)}")
print(f"sql_bytes={SQL.stat().st_size}")

