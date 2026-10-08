from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "智能驾驶汽车用户使用体验调研问卷_1008.docx"
OUTPUT = ROOT / "智能驾驶汽车用户使用体验调研问卷_规整及系统设计_1008.docx"
ASSETS = ROOT / ".work" / "assets"
EXPECTED_SHA256 = "0B9301A49A1BE69AD891706A4D4BEA1764D8431663C72FCA7F8BA85B396EA695"
USER_URL = "https://ads-experience-survey-ui.shchen-lmars.chatgpt.site/"
ADMIN_URL = "https://ads-experience-survey-ui.shchen-lmars.chatgpt.site/admin.html"


def file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def restore_package_parts(source: Path, target: Path, part_names: list[str]) -> None:
    """Restore untouched template parts byte-for-byte after python-docx saves."""
    temp = target.with_suffix(".repacked.docx")
    with zipfile.ZipFile(source, "r") as source_zip, zipfile.ZipFile(target, "r") as target_zip:
        source_bytes = {name: source_zip.read(name) for name in part_names}
        with zipfile.ZipFile(temp, "w") as output_zip:
            for info in target_zip.infolist():
                payload = source_bytes.get(info.filename, target_zip.read(info.filename))
                output_zip.writestr(info, payload)
    temp.replace(target)


def set_run_font(run, name: str = "Noto Sans SC", size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "D9D9D9", size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def style_table(table, widths: list[float], header_fill: str = "183B66") -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    for row_idx, row in enumerate(table.rows):
        keep_row_together(row)
        for col_idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "F4F8FB")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.1
                if col_idx == 0 and row_idx > 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(run, size=9.0, bold=(row_idx == 0), color="FFFFFF" if row_idx == 0 else "172033")


def set_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_figure(document: Document, image_path: Path, caption: str, width: float, alt: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    shape = paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    set_alt_text(shape, caption, alt)
    caption_p = document.add_paragraph(style="SurveyNote")
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_before = Pt(0)
    caption_p.paragraph_format.space_after = Pt(10)
    run = caption_p.add_run(caption)
    set_run_font(run, size=9.0, color="4A5968")


def add_figure_before(paragraph, image_path: Path, caption: str, width: float, alt: str) -> None:
    pic_paragraph = paragraph.insert_paragraph_before()
    pic_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_paragraph.paragraph_format.space_before = Pt(8)
    pic_paragraph.paragraph_format.space_after = Pt(4)
    shape = pic_paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    set_alt_text(shape, caption, alt)
    paragraph.text = caption
    paragraph.style = "SurveyNote"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(10)
    for run in paragraph.runs:
        set_run_font(run, size=9.0, color="4A5968")


def replace_prefix_in_runs(paragraph, replacement: str) -> None:
    full = paragraph.text
    match = re.match(r"^Q\d+", full)
    if not match:
        raise ValueError(f"Question prefix missing: {full!r}")
    old = match.group(0)
    remaining = len(old)
    replacement_pending = replacement
    for run in paragraph.runs:
        if remaining <= 0:
            break
        take = min(remaining, len(run.text))
        if take:
            suffix = run.text[take:]
            run.text = replacement_pending + suffix
            replacement_pending = ""
            remaining -= take
    if remaining:
        raise ValueError(f"Question prefix spans unexpected runs: {full!r}")


def add_hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "176F8C")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_paragraph(paragraph, text: str, style: str | None = None, bold_lead: str | None = None) -> None:
    paragraph.clear()
    if style:
        paragraph.style = style
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)


def add_subheading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="SurveySubtitle")
    paragraph.paragraph_format.space_before = Pt(11)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run_font(run, size=12.0, bold=True, color="000000")


def set_update_fields(document: Document) -> None:
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main() -> None:
    if file_sha256(SOURCE) != EXPECTED_SHA256:
        raise SystemExit("Source DOCX SHA-256 mismatch; refusing to edit a changed reference.")
    required_assets = [
        ASSETS / "figure_01_region_distribution.png",
        ASSETS / "figure_02_cumulative_trend.png",
        ASSETS / "figure_03_brand_ranking.png",
        ASSETS / "figure_04_database_er.png",
    ]
    for asset in required_assets:
        if not asset.exists():
            raise SystemExit(f"Missing asset: {asset}")

    document = Document(str(SOURCE))
    document.core_properties.title = "智能驾驶汽车用户调研问卷及系统设计"
    document.core_properties.subject = "规整问卷 UI界面与数据库设计"
    document.core_properties.keywords = "智能驾驶 问卷 UI 数据库 MySQL"

    # Use the existing report-summary slots for the three requested simulations.
    by_text = {p.text.strip(): p for p in document.paragraphs if p.text.strip()}
    add_figure_before(
        by_text["【留一幅中国地图，划分省、市的区域，每个区域显示当前有多少数据】"],
        required_assets[0],
        "图 1 地域分布模拟（非行政边界图，示例数据）",
        6.85,
        "区域气泡分布模拟，显示华北、华东、华南、华中、西南、西北、东北的示例答卷数量。",
    )
    add_figure_before(
        by_text["【画一幅统计图，横轴是时间，精确到月份，纵轴是累积的记录总数。】"],
        required_assets[1],
        "图 2 月度累计记录模拟（示例数据）",
        6.85,
        "按月份显示累计答卷数量的折线面积图示例。",
    )
    add_figure_before(
        by_text["【统计25个主要品牌类型比例，从高往低排列，折线图】"],
        required_assets[2],
        "图 3 主要品牌占比排序模拟（25 类，示例数据）",
        6.95,
        "25个匿名品牌按答卷占比从高到低排列的折线图示例。",
    )

    # Renumber every real question, preserving its existing run-level formatting.
    questions = [p for p in document.paragraphs if p.style.name == "SurveyQuestion"]
    if len(questions) != 61:
        raise SystemExit(f"Expected 61 SurveyQuestion paragraphs, found {len(questions)}")
    for index, paragraph in enumerate(questions, start=1):
        replace_prefix_in_runs(paragraph, f"Q{index:02d}")

    # Small wording and formatting cleanups while preserving meaning.
    question_by_code = {re.match(r"^(Q\d+)", p.text).group(1): p for p in questions}
    set_paragraph(question_by_code["Q04"], "Q04  您长期居住的地区是？（仅填写省、市）", "SurveyQuestion")
    set_paragraph(question_by_code["Q49"], "Q49  您对智驾行业的哪些问题感到困惑？最多选3项。〔全部合格受访者｜最多选3项〕", "SurveyQuestion")

    paragraphs = document.paragraphs
    heading_updates = {
        "2. 个人背景与驾驶经验  〔Q04–Q11〕": "1. 个人背景与驾驶经验  〔Q01–Q06〕",
        "3. 当前车辆与购车情况  〔Q12–Q18〕": "2. 当前车辆与购车情况  〔Q07–Q10〕",
        "4. 智驾认知、培训与使用  〔Q19–Q30〕": "3. 智驾认知、培训与使用  〔Q11–Q14〕",
        "5. 功能需求与品牌偏好  〔Q31–Q35〕": "4. 功能需求与品牌偏好  〔Q15〕",
        "6. AI赋能路径与效果  〔Q36–Q42〕": "5. AI赋能路径与效果  〔Q16–Q20〕",
        "7. 功能障碍与风险事件  〔Q43–Q60〕": "6. 功能障碍与风险事件  〔Q21–Q38〕",
        "8. 安全、信任与购买意向  〔Q61–Q66〕": "7. 安全、信任与购买意向  〔Q39–Q43〕",
        "9. 全生命周期价值  〔Q67–Q73〕": "8. 全生命周期价值  〔Q44–Q48〕",
        "10. 政府监管、行业治理与政策期待  〔Q74–Q80〕": "9. 政府监管、行业治理与政策期待  〔Q49–Q56〕",
        "11. 总体评价与开放意见  〔Q81–Q84〕": "10. 总体评价与开放意见  〔Q57–Q59〕",
        "12. 自愿回访（与主答卷分表保存）  〔Q85–Q86〕": "11. 自愿回访（与主答卷分表保存）  〔Q60–Q61〕",
    }
    for paragraph in paragraphs:
        normalized = paragraph.text.strip()
        if normalized in heading_updates:
            set_paragraph(paragraph, heading_updates[normalized], "Heading 1")

    exact_replacements = {
        "是否需要调研报告：邮箱？": "如需接收调研报告，可在自愿回访模块选择邮件联系。",
        "说明：不收集详细地址。__省      市 ______________": "说明：仅统计到省、市，不收集详细地址。",
        "说明：不要填写车牌、姓名、精确地址等可识别信息。___最好能够支持语音输入，并且转为文字。____________": "说明：不要填写车牌、姓名、手机号、精确地址等可识别信息；线上版本可支持语音转文字后由用户确认。",
        "说明：可填写昵称或称呼，不要求真实姓名：________________": "说明：若愿意回访，请填写昵称或称呼及电话或邮箱；联系方式与主答卷分表保存。",
    }
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text in exact_replacements:
            set_paragraph(paragraph, exact_replacements[text], paragraph.style.name)

    # Fix a duplicated checkbox and normalize the formerly loose Q49 option list.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if "□□ 短期租用" in cell.text:
                    cell.text = cell.text.replace("□□ 短期租用", "□ 短期租用")
    option_replacements = {
        "宣传乱象": "□ 宣传信息或功能命名混乱",
        "发生争议后，数据和证据难以获得；": "□ 发生争议后，数据和证据难以获得",
        "电池容易引起火灾；": "□ 动力电池安全风险",
        "电池衰减": "□ 电池性能衰减",
        "智驾系统功能边界模糊；": "□ 智驾系统功能边界模糊",
        "其他：？": "□ 其他：________________",
    }
    for paragraph in document.paragraphs:
        if paragraph.text.strip() in option_replacements:
            set_paragraph(paragraph, option_replacements[paragraph.text.strip()], "SurveyOption")

    # Replace the terse role notes with a complete UI design recommendation.
    role_start = next(i for i, p in enumerate(document.paragraphs) if p.text.strip() == "普通用户:")
    role_end = next(i for i, p in enumerate(document.paragraphs) if p.text.strip() == "四、数据库设计")
    role_paragraphs = document.paragraphs[role_start:role_end]
    ui_content = [
        ("普通用户端", "普通用户端：采用分步答卷、清晰进度、矩阵题横向滚动、自动保存提示和隐私提醒；手机端将侧栏折叠为顶部进度，操作按钮固定在底部。"),
        ("普通用户端入口", f"普通用户端入口：{USER_URL}"),
        ("管理员端", "管理员端：采用高密度数据工作台，包含累计提交、有效答卷、风险事件、待复核记录、趋势图、筛选表格、记录详情、补充字段和脱敏导出。"),
        ("管理员端入口", f"管理员端入口：{ADMIN_URL}"),
        ("响应式策略", "响应式策略：电脑网页端使用侧栏加主工作区；手机网页端使用单列布局、可横向滚动矩阵与全屏记录抽屉，保证触控目标与正文可读。"),
        ("推荐技术", "推荐技术：生产版本建议使用 TypeScript + React，并配合 Tailwind CSS 与成熟的无障碍组件库。TypeScript 负责数据结构与接口约束，React 负责复杂表单和管理台状态，Tailwind CSS 负责响应式设计令牌与交付效率。"),
        ("推荐依据", "推荐依据：截至2026年，GitHub Octoverse 2025显示TypeScript已成为GitHub使用量最高的语言；Stack Overflow 2026开发者调查中React使用率为41.5%，居前端框架前列；React官方文档说明主流生产级React框架均支持TypeScript。"),
        ("原型说明", "原型说明：本次链接展示UI与主要交互，使用模拟数据，不连接真实账号、主答卷或联系方式数据库。"),
    ]
    for paragraph, (lead, text) in zip(role_paragraphs, ui_content):
        set_paragraph(paragraph, text, "ReportBody", bold_lead=f"{lead}：" if text.startswith(f"{lead}：") else None)
    for paragraph in role_paragraphs[len(ui_content):]:
        paragraph.clear()

    # Add source links under the UI recommendation without turning the document into a bibliography-heavy report.
    heading_db = next(p for p in document.paragraphs if p.text.strip() == "四、数据库设计")
    sources_p = heading_db.insert_paragraph_before()
    sources_p.style = "SurveyNote"
    lead = sources_p.add_run("技术依据链接：")
    set_run_font(lead, bold=True, size=9.0)
    add_hyperlink(sources_p, "https://github.blog/news-insights/octoverse/octoverse-a-new-developer-joins-github-every-second-as-ai-leads-typescript-to-1/", "GitHub Octoverse 2025")
    sources_p.add_run("；")
    add_hyperlink(sources_p, "https://survey.stackoverflow.co/2026/technology/data/web-framework", "Stack Overflow 2026 调查")
    sources_p.add_run("；")
    add_hyperlink(sources_p, "https://react.dev/learn/typescript", "React TypeScript 文档")

    # Give the UI and database design their own clean page starts.
    heading_roles = next(p for p in document.paragraphs if p.text.strip() == "三、用户与角色体系及对应功能")
    heading_roles.paragraph_format.page_break_before = True
    heading_db.paragraph_format.page_break_before = True

    # Populate the database section from the existing trailing slots.
    trailing = document.paragraphs[-4:]
    set_paragraph(trailing[0], "数据库采用MySQL 8.0逻辑模型，将问卷定义、版本、题目、匿名答卷、风险事件、回访联系方式、扩展字段和审计任务分域管理。", "ReportBody")
    set_paragraph(trailing[1], "主答卷不直接保存姓名、手机号或邮箱；回访联系方式单独建表并保存密文。问卷发布后以版本锁定题目，确保历史答卷可追溯。", "ReportBody")
    set_paragraph(trailing[2], "随附管理脚本：问卷调查系统_数据库管理_MySQL8.sql。脚本包含建库、主外键、索引、校验约束、脱敏视图、初始角色和常用管理查询。", "ReportBody")
    trailing[3].clear()

    add_figure(
        document,
        required_assets[3],
        "图 4 问卷调查系统数据库逻辑结构",
        7.05,
        "数据库逻辑关系图，展示账户权限、问卷定义、题目结构、匿名答卷、答案事件、回访联系方式、字段扩展与审计导出。",
    )

    add_subheading(document, "数据库表设计")
    table_rows = [
        ("app_user", "系统账户", "login_name、password_hash、account_status；不存主答卷"),
        ("app_role", "角色定义", "ADMIN、ANALYST、AUDITOR"),
        ("app_user_role", "账户角色关系", "user_id + role_id 联合主键"),
        ("survey", "问卷主表", "survey_code、survey_status、owner_user_id"),
        ("survey_version", "问卷版本", "version_no、published_at；已发布版本只读"),
        ("survey_section", "章节与分支", "display_order、visibility_rule_json"),
        ("survey_question", "题目定义", "question_code、question_type、validation_json、branching_json"),
        ("question_option", "题目选项", "option_code、option_value、is_exclusive"),
        ("survey_response", "匿名主答卷", "respondent_key、状态、地区、风险等级、5个预留VARCHAR字段"),
        ("survey_answer", "答案记录", "answer_text、answer_number、answer_json；每答卷每题唯一"),
        ("incident_detail", "风险事件", "功能、严重度、提示、解决状态、开放描述"),
        ("followup_contact", "自愿回访", "手机号和邮箱密文、同意状态、用途、期限；与主答卷分表"),
        ("custom_field_definition", "扩展字段定义", "字段编码、数据类型、启用状态"),
        ("response_custom_field", "扩展字段值", "按类型分列保存，记录编辑人和更新时间"),
        ("export_task", "脱敏导出任务", "过滤条件、格式、状态、过期时间"),
        ("admin_audit_log", "管理员审计", "操作人、对象、前后值、请求编号；只追加"),
    ]
    table = document.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "数据表"
    table.rows[0].cells[1].text = "用途"
    table.rows[0].cells[2].text = "关键字段与关系"
    for name, purpose, key_fields in table_rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = name, purpose, key_fields
    style_table(table, [1.55, 1.35, 4.15])

    add_subheading(document, "预留字符字段")
    p = document.add_paragraph(style="ReportBody")
    p.add_run("在 survey_response 表预留 ")
    r = p.add_run("reserved_char_01 至 reserved_char_05")
    set_run_font(r, bold=True)
    p.add_run(" 五个 VARCHAR(255) 可空字段，满足近期小范围调整；长期或多类型扩展使用 custom_field_definition 与 response_custom_field，避免无限增加主表列。")
    reserved = document.add_table(rows=1, cols=3)
    reserved.rows[0].cells[0].text = "字段"
    reserved.rows[0].cells[1].text = "类型"
    reserved.rows[0].cells[2].text = "使用约束"
    for idx in range(1, 6):
        cells = reserved.add_row().cells
        cells[0].text = f"reserved_char_{idx:02d}"
        cells[1].text = "VARCHAR(255) NULL"
        cells[2].text = "启用前登记字段含义、责任人和生效版本；不得保存口令或明文联系方式"
    style_table(reserved, [2.1, 1.8, 3.15], header_fill="285783")

    add_subheading(document, "数据管理与安全说明")
    notes = [
        "版本与追溯：问卷发布后新建版本，不原位修改历史题目；答卷始终关联 survey_version_id。",
        "最小化与分离：主答卷保存匿名 respondent_key；手机号和邮箱仅在自愿回访时采集并进入 followup_contact 密文列。",
        "权限与审计：管理员、分析员、审计员分权；记录编辑、字段变更和导出均写入 admin_audit_log。",
        "性能与导出：按提交时间、状态、地区和风险等级建组合索引；大批量导出通过 export_task 异步执行并设置失效时间。",
        "清理与撤回：主表采用软删除；回访同意可撤回，到期后应清除或不可逆匿名化联系方式。",
    ]
    for note in notes:
        paragraph = document.add_paragraph(style="ReportBody")
        paragraph.style = "ReportBody"
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        run = paragraph.add_run("• " + note)
        set_run_font(run)

    set_update_fields(document)
    document.save(str(OUTPUT))

    restore_package_parts(
        SOURCE,
        OUTPUT,
        [
            "word/header1.xml",
            "word/footer1.xml",
            "word/theme/theme1.xml",
            "word/numbering.xml",
            "word/styles.xml",
            "word/webSettings.xml",
            "word/fontTable.xml",
            "word/footnotes.xml",
            "word/endnotes.xml",
            "customXml/item1.xml",
            "customXml/itemProps1.xml",
            "customXml/_rels/item1.xml.rels",
        ],
    )

    if file_sha256(SOURCE) != EXPECTED_SHA256:
        raise SystemExit("Source DOCX changed during authoring")
    print(OUTPUT)
    print(f"size={OUTPUT.stat().st_size}")
    print(f"sha256={file_sha256(OUTPUT)}")


if __name__ == "__main__":
    main()

